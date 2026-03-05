# -*- coding: utf-8 -*-
"""
مسح بيانات التدريب (متدربين-تسجيلات، مدربين، دورات، دفعات، قاعات) ثم استيراد من academy_requested_clean_sheets.xlsx
شغّل من مجلد المشروع: python clear_and_import_academy.py
لا يعدّل قواعد البيانات ولا الشاشات — يستخدم الجداول والحقول الحالية فقط.
"""
import os
import sys
import json
import re
import pyodbc
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("ثبّت openpyxl: pip install openpyxl")
    sys.exit(1)


def get_connection():
    base = os.path.dirname(os.path.abspath(__file__))
    config_path = os.path.join(base, 'db_config.json')
    if not os.path.exists(config_path):
        print('لا يوجد db_config.json')
        sys.exit(1)
    with open(config_path, 'r', encoding='utf-8') as f:
        config = json.load(f)
    server = config.get('server', '.')
    port = config.get('port', '1433')
    database = config.get('database', 'Place2026DB')
    username = config.get('username', '')
    password = config.get('password', '')
    use_trusted = config.get('use_trusted', False)
    if use_trusted:
        conn_str = f'DRIVER={{ODBC Driver 17 for SQL Server}};SERVER={server},{port};DATABASE={database};Trusted_Connection=yes;'
    else:
        conn_str = f'DRIVER={{ODBC Driver 17 for SQL Server}};SERVER={server},{port};DATABASE={database};UID={username};PWD={password};'
    return pyodbc.connect(conn_str, timeout=15)


def run_delete(cursor):
    """مسح بالترتيب (احترام المفاتيح الأجنبية)."""
    order = [
        'StudentPayments',
        'Attendance',
        'WeeklyExams',
        'WeeklyProgress',
        'TrainingPlans',
        'Enrollments',
        'CourseBatches',
        'TrainingOffers',
        'Courses',
        'Trainers',
        'Classrooms',
    ]
    for table in order:
        try:
            cursor.execute(f'DELETE FROM [{table}]')
            print(f'  حذف: {table} ({cursor.rowcount} صف)')
        except pyodbc.Error as e:
            if 'Invalid object name' in str(e) or 'does not exist' in str(e).lower():
                print(f'  تخطي (الجدول غير موجود): {table}')
            else:
                raise


def safe_str(val, max_len=200):
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() in ('none', 'nan', '#n/a'):
        return None
    return s[:max_len] if len(s) > max_len else s


def safe_int(val):
    if val is None:
        return None
    try:
        return int(float(val))
    except (TypeError, ValueError):
        return None


def safe_decimal(val):
    if val is None:
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(base, 'academy_requested_clean_sheets.xlsx')
    if not os.path.exists(excel_path):
        print('لا يوجد ملف academy_requested_clean_sheets.xlsx في مجلد المشروع.')
        sys.exit(1)

    print('الاتصال بقاعدة البيانات...')
    conn = get_connection()
    conn.autocommit = False
    cursor = conn.cursor()

    try:
        print('\n--- 1) مسح بيانات التدريب ---')
        run_delete(cursor)
        conn.commit()

        print('\n--- 2) استيراد من الملف ---')
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

        # --- القاعات ---
        print('  قراءة ورقة القاعات...')
        ws_rooms = wb['بيانات_القاعات']
        rooms_seen = set()
        room_name_col = 0  # room_name
        for row in ws_rooms.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            name = safe_str(row[room_name_col], 50)
            if not name or name.lower() in ('n/a', 'na', ''):
                continue
            norm = name.strip().lower()
            if norm in rooms_seen:
                continue
            rooms_seen.add(norm)
            try:
                cursor.execute(
                    'INSERT INTO Classrooms (RoomName, Capacity, IsActive) VALUES (?, 20, 1)',
                    (name,)
                )
            except pyodbc.Error as e:
                print(f'  تحذير قاعة {name}: {e}')
        print(f'  تم إدراج {len(rooms_seen)} قاعة.')
        conn.commit()

        # --- المدربون ---
        ws_trainers = wb['بيانات_المدربين']
        trainers_seen = set()
        for row in ws_trainers.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            name = safe_str(row[0], 100)  # trainer_name
            if not name:
                continue
            if re.match(r'^\d+$', str(name).strip()):
                continue
            norm = name.strip().lower()
            if norm in trainers_seen:
                continue
            trainers_seen.add(norm)
            try:
                cursor.execute(
                    'INSERT INTO Trainers (FullName, Specialization, Phone) VALUES (?, NULL, NULL)',
                    (name.strip(),)
                )
            except pyodbc.Error as e:
                print(f'  تحذير مدرب {name}: {e}')
        print(f'  تم إدراج {len(trainers_seen)} مدرب.')
        conn.commit()

        # --- دورة واحدة ثم الدفعات ---
        cursor.execute(
            "INSERT INTO Courses (CourseName, LevelOrder, DefaultPrice, Description) VALUES (N'دورة لغة إنجليزية - أكاديمي', 1, 0, NULL)"
        )
        cursor.execute('SELECT SCOPE_IDENTITY()')
        course_id = int(cursor.fetchone()[0])
        print('  تم إنشاء دورة واحدة (CourseID=%s)' % course_id)

        ws_batches = wb['بيانات_الدفعات']
        batch_names_seen = set()
        for row in ws_batches.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            batch_name = safe_str(row[0], 100)  # batch
            if not batch_name:
                continue
            norm = batch_name.strip().lower()
            if norm in batch_names_seen:
                continue
            batch_names_seen.add(norm)
            try:
                cursor.execute(
                    """INSERT INTO CourseBatches (CourseID, TrainerID, RoomID, BatchName, StartDate, EndDate, ScheduleDescription, Status)
                       VALUES (?, NULL, NULL, ?, NULL, NULL, NULL, 'Active')""",
                    (course_id, batch_name.strip())
                )
            except pyodbc.Error as e:
                print(f'  تحذير دفعة {batch_name}: {e}')
        print(f'  تم إدراج {len(batch_names_seen)} دفعة.')
        conn.commit()

        # --- بناء خريطة اسم الدفعة -> BatchID ---
        cursor.execute('SELECT BatchID, BatchName FROM CourseBatches')
        batch_name_to_id = {}
        for r in cursor.fetchall():
            batch_name_to_id[(r[1] or '').strip().lower()] = r[0]
            batch_name_to_id[(r[1] or '').strip()] = r[0]

        # --- الطلاب: مرشحون + تسجيلات ---
        ws_students = wb['بيانات_الطلاب']
        headers = [c.value for c in next(ws_students.iter_rows(min_row=1, max_row=1))]
        try:
            i_name = headers.index('student_name')
            i_phone = headers.index('phone_digits') if 'phone_digits' in headers else headers.index('phone')
        except ValueError:
            i_name = 0
            i_phone = 1
        i_email = headers.index('email') if 'email' in headers else -1
        i_batch = headers.index('batch') if 'batch' in headers else -1

        phone_to_candidate_id = {}
        cursor.execute('SELECT CandidateID, Phone FROM Candidates WHERE Phone IS NOT NULL AND LEN(RTRIM(Phone))>0')
        for r in cursor.fetchall():
            if r[1]:
                phone_to_candidate_id[str(r[1]).strip()] = r[0]
                phone_to_candidate_id[str(r[1]).replace(' ', '')] = r[0]

        inserted_candidates = 0
        inserted_enrollments = 0
        skipped_enrollment = 0

        for row in ws_students.iter_rows(min_row=2, values_only=True):
            if not row or len(row) <= max(i_name, i_phone):
                continue
            full_name = safe_str(row[i_name], 100)
            phone = safe_str(row[i_phone], 50) if i_phone >= 0 else None
            if not full_name:
                continue
            email = safe_str(row[i_email], 100) if i_email >= 0 and i_email < len(row) else None
            batch_name = safe_str(row[i_batch], 100) if i_batch >= 0 and i_batch < len(row) else None

            phone_key = (phone or '').replace(' ', '').strip()
            if not phone_key and email:
                phone_key = (email or '').strip()

            cand_id = None
            if phone_key and phone_key in phone_to_candidate_id:
                cand_id = phone_to_candidate_id[phone_key]
            else:
                try:
                    cursor.execute(
                        """INSERT INTO Candidates (FullName, Phone, Email, Status, InterestLevel)
                           VALUES (?, ?, ?, 'TrainingOnly', 'High')""",
                        (full_name, phone or None, email)
                    )
                    cursor.execute('SELECT SCOPE_IDENTITY()')
                    cand_id = int(cursor.fetchone()[0])
                    if phone_key:
                        phone_to_candidate_id[phone_key] = cand_id
                    inserted_candidates += 1
                except pyodbc.Error as e:
                    print(f'  تحذير مرشح {full_name}: {e}')
                    continue

            if not batch_name or not cand_id:
                continue
            batch_key = batch_name.strip().lower()
            batch_id = batch_name_to_id.get(batch_key) or batch_name_to_id.get(batch_name.strip())
            if not batch_id:
                skipped_enrollment += 1
                continue
            try:
                cursor.execute(
                    """INSERT INTO Enrollments (BatchID, CandidateID, EnrollmentDate, Status, AgreedPrice)
                       VALUES (?, ?, GETDATE(), 'Active', 0)""",
                    (batch_id, cand_id)
                )
                inserted_enrollments += 1
            except pyodbc.Error as e:
                skipped_enrollment += 1

        print(f'  مرشحون جدد: {inserted_candidates} | تسجيلات (Enrollments): {inserted_enrollments} | تخطي (دفعة غير موجودة): {skipped_enrollment}')
        conn.commit()

        wb.close()
        print('\nتم المسح والاستيراد بنجاح.')
    except Exception as e:
        conn.rollback()
        print('خطأ:', e)
        raise
    finally:
        cursor.close()
        conn.close()


if __name__ == '__main__':
    main()
