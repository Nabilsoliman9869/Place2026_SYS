# -*- coding: utf-8 -*-
"""تصدير خريطة اختبار تكامل النظام إلى Excel — قسمان منفصلان (توظيف / تدريب)، نافذة + مستخدم لكل عملية.
اليوزرات من app.py و create_users (قاعدة البيانات)."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

OUTPUT = "خريطة_اختبار_تكامل_النظام.xlsx"
# إذا ظهر Permission denied: أغلق الملف من Excel ثم شغّل السكربت مرة أخرى

def header_style():
    return Font(bold=True)

def sheet_header(ws, title, start_row=1):
    ws.cell(start_row, 1, title).font = Font(bold=True, size=14)
    return start_row + 2

def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

    # --- 0. توضيح القسمين (مبيعات التوظيف | مبيعات التدريب | نقطة التكامل) ---
    ws0 = wb.create_sheet("0_توضيح_القسمين", 0)
    r = 1
    ws0.cell(r, 1, "خريطة الاختبار — قسمان منفصلان يتكاملان عند فشل المترشح لوظيفة").font = Font(bold=True, size=14)
    r += 2
    ws0.cell(r, 1, "قسم التوظيف (منفصل)").font = Font(bold=True, size=12)
    r += 1
    ws0.cell(r, 1, "• مبيعات التوظيف + حملاتهم: استقطاب مهتمين بالوظيفة.")
    r += 1
    ws0.cell(r, 1, "• تعريف المهتم لأول مرة = تسجيل ليد/مهتم بالوظيفة (استقطاب للتوظيف).")
    r += 1
    ws0.cell(r, 1, "• المسار: مهتم → مرشح → تقييم → مقبول → ترشيح → مقابلة → تعيين.")
    r += 3
    ws0.cell(r, 1, "قسم التدريب (منفصل) — أدوار مبيعات التدريب تحاكي أدوار مبيعات التوظيف").font = Font(bold=True, size=12)
    r += 1
    ws0.cell(r, 1, "• مبيعات التدريب: يعمل بوست (حملة)، يسجّل مهتم تدريب، يختبره مبدئياً، يحجز له موعد اختبار مواهب تدريب، يتابع الموعد حتى يصل المهتم، يحصل منه رسوم الاشتراك.")
    r += 1
    ws0.cell(r, 1, "• يحدد مختبر المواهب للتدريب (TA-Training). بعد دفع رسم الامتحان ورغبة المتدرب يأخذه منسق التدريب (تسجيل في الدفعة).")
    r += 1
    ws0.cell(r, 1, "• المسار: بوست → تسجيل مهتم تدريب → اختبار مبدئي → حجز موعد اختبار مواهب تدريب → متابعة → رسوم اشتراك → تحديد مختبر مواهب تدريب → دفع رسم الامتحان + رغبة المتدرب → منسق التدريب يُلحقه بالدفعة → حضور → تخرج.")
    r += 3
    ws0.cell(r, 1, "نقطة التكامل الوحيدة").font = Font(bold=True, size=12)
    r += 1
    ws0.cell(r, 1, "عندما يفشل المترشح في مسار الوظيفة (مرفوض أو غير مناسب) → يُحوّل إلى التدريب → مبيعات التدريب (أو المنسق) تتصل به وتُسوّق له الاشتراك في التدريب (دورة / دفعة).")
    ws0.column_dimensions['A'].width = 95

    # --- 1. مسار التوظيف (مع النافذة + المستخدم) ---
    ws1 = wb.create_sheet("1_مسار_التوظيف", 1)
    r = sheet_header(ws1, "١ — مسار التوظيف (من الاستقطاب إلى التعيين)")
    headers = ["#", "الخطوة", "النافذة", "المستخدم", "القسم/الدور", "الهدف"]
    for c, h in enumerate(headers, 1):
        cell = ws1.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    # ترتيب، خطوة، نافذة، مستخدم، قسم، هدف
    data1 = [
        (1, "استقطاب مهتم (تعريف المهتم بالوظيفة)", "الورقة اليومية / Daily Sheet أو استيراد Leads", "مدير تسويق حسب الإعداد", "التسويق", "ظهور المهتم في لوحة المبيعات"),
        (2, "تحويل إلى مرشح", "لوحة المبيعات — تحويل Lead إلى Candidate", "sales", "المبيعات", "إنشاء مرشح وربطه بالحملة"),
        (3, "فرز مبدئي", "Workbench — Feedback (Valid/Invalid)", "recruiter", "التوظيف", "تغيير الحالة إلى Talent_Pool أو توجيه للتدريب"),
        (4, "حجز اختبار", "Scheduling — Book Test", "recruiter", "التوظيف", "ربط المرشح بموعد الاختبار"),
        (5, "إجراء التقييم", "Conduct Tests — Evaluate", "ta_rec أو مستخدم Talent_Recruitment", "المواهب (توظيف)", "إدخال CEFR والقرار (مقبول/مرفوض)"),
        (6, "ترشيح لطلب عميل", "Matching — Confirm Match", "alloc_sp أو alloc_mgr", "الترشيح", "ربط المرشح بطلب توظيف عميل"),
        (7, "مقابلة العميل", "Interviews / Interview Feedback", "account أو recruiter", "مدير الحسابات/التوظيف", "تسجيل نتيجة المقابلة"),
        (8, "تعيين", "Hiring & Onboarding — Finalize", "recruiter", "التوظيف", "تاريخ البدء، راتب، نوع عقد، إغلاق الملف"),
    ]
    for row in data1:
        r += 1
        for c, val in enumerate(row, 1):
            ws1.cell(r, c, val).border = border
    for col in ['A','B','C','D','E','F']:
        ws1.column_dimensions[col].width = 18 if col in ['A','D'] else 22 if col == 'C' else 28 if col == 'B' else 35 if col == 'E' else 42

    # --- 2. مسار مبيعات التدريب (محاكاة مبيعات التوظيف) — خطوة بخطوة ---
    ws2 = wb.create_sheet("2_مسار_مبيعات_التدريب", 2)
    r = sheet_header(ws2, "٢ — مسار مبيعات التدريب (تحاكي مبيعات التوظيف)")
    ws2.cell(r, 1, "أدوار مبيعات التدريب = نفس منطق مبيعات التوظيف: بوست → تسجيل مهتم → اختبار مبدئي → حجز موعد → متابعة → رسوم → مختبر مواهب تدريب → بعد الدفع ورغبة المتدرب يأخذه المنسق.").font = Font(italic=True)
    r += 2
    headers = ["#", "الخطوة", "النافذة", "المستخدم", "الهدف"]
    for c, h in enumerate(headers, 1):
        cell = ws2.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data2 = [
        (1, "يعمل بوست (حملة تدريب)", "حملات التدريب / بوست", "train_sales", "التسويق للدورات وجذب مهتمين"),
        (2, "يسجّل مهتم تدريب (تعريف المهتم لأول مرة)", "تسجيل مهتم بالتدريب", "train_sales", "إنشاء سجل مهتم بالدورة/التدريب"),
        (3, "يختبره مبدئياً", "اختبار مبدئي / فرز مهتمي التدريب", "train_sales", "تحديد صلاحية للمتابعة"),
        (4, "يحجز له موعد اختبار مواهب تدريب", "جدولة مواعيد اختبار التدريب", "train_sales", "ربط المهتم بموعد اختبار TA-Training"),
        (5, "يتابع الموعد حتى يصل المهتم", "متابعة المواعيد / الحضور", "train_sales", "التأكد من حضور المهتم للموعد"),
        (6, "يحصل منه قيمة رسوم الاشتراك", "استلام رسوم الاشتراك", "train_sales", "تسجيل استلام رسوم الاشتراك"),
        (7, "يحدد مختبر المواهب للتدريب", "Talent Evaluate — إجراء التقييم (نموذج: C,F,P,G,V سطر واحد؛ لا Rejected)", "ta_train (مختبر مواهب التدريب)", "إجراء اختبار المواهب للتدريب — في التدريب لا يُرفض أحد"),
        (8, "بعد دفع رسم الامتحان + رغبة المتدرب → منسق التدريب يأخذه", "تسجيل في الدفعة (Academy — تسجيل طلاب)", "train_coord", "إلحاق المتدرب بالدفعة بعد الدفع ورغبته"),
    ]
    for row in data2:
        r += 1
        for c, val in enumerate(row, 1):
            ws2.cell(r, c, val).border = border
    ws2.column_dimensions['A'].width = 6
    ws2.column_dimensions['B'].width = 48
    ws2.column_dimensions['C'].width = 42
    ws2.column_dimensions['D'].width = 28
    ws2.column_dimensions['E'].width = 48

    # --- 2ب. مسار التدريب (بعد الإلحاق: حضور، تخرج) ---
    ws2b = wb.create_sheet("2ب_بعد_الإلحاق", 3)
    r = sheet_header(ws2b, "٢ب — بعد إلحاق المتدرب بالدفعة (حضور، تخرج)")
    headers = ["#", "الخطوة", "النافذة", "المستخدم", "الهدف"]
    for c, h in enumerate(headers, 1):
        cell = ws2b.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data2b = [
        (1, "إعداد الدورات والدفعات", "Academy — إضافة دورة، إضافة دفعة (Wave)", "train_mgr / train_head / train_lead / train_coord", "إنشاء دورة ودفعة وتواريخ وقاعة ومدرب"),
        (2, "تسجيل طلاب في الدفعة (المنسق بعد الدفع ورغبة المتدرب)", "Academy — تسجيل طلاب (batch_details)", "train_coord أو trainer1", "إلحاق متدرب بالدفعة"),
        (3, "تسجيل الحضور", "Attendance — مصفوفة الحضور", "trainer1", "تسجيل حضور/غياب لكل يوم وجلسة"),
        (4, "المدفوعات والفوترة", "Student Finance — Add Payment / Print Invoice", "مدير مالية أو من شاشة الطالب", "تسجيل الدفعات وطباعة الفواتير"),
        (5, "التخرج ومراجعة الخريجين", "تفاصيل الدفعة — Graduate", "trainer1 أو manager", "تخريج الطالب والموافقة"),
        (6, "قرار الطالب بعد التخرج", "Student Decision", "من له صلاحية", "رغبة الخريج (عودة للتوظيف / دورة أخرى)"),
    ]
    for row in data2b:
        r += 1
        for c, val in enumerate(row, 1):
            ws2b.cell(r, c, val).border = border
    ws2b.column_dimensions['A'].width = 6
    ws2b.column_dimensions['B'].width = 50
    ws2b.column_dimensions['C'].width = 48
    ws2b.column_dimensions['D'].width = 38
    ws2b.column_dimensions['E'].width = 48

    # --- 2ج. تفاصيل نموذج تقييم التدريب (Talent Evaluate) ---
    ws2c = wb.create_sheet("2ج_نموذج_تقييم_التدريب", 4)
    r = sheet_header(ws2c, "٢ج — نموذج تقييم مواهب التدريب (Talent Evaluate) — آخر تحديث")
    ws2c.cell(r, 1, "البند").font = header_style()
    ws2c.cell(r, 2, "الوصف").font = header_style()
    r += 1
    ws2c.cell(r, 1, "حقول الدرجات C, F, P, G, V").border = border
    ws2c.cell(r, 2, "على نفس السطر (سطر واحد) لسهولة الاستخدام").border = border
    r += 1
    ws2c.cell(r, 1, "القرار النهائي (Final Recommendation)").border = border
    ws2c.cell(r, 2, "في تقييم التدريب: لا يظهر خيار Rejected — الخياران فقط Accepted أو Training (في التدريب لا يُرفض أحد)").border = border
    ws2c.column_dimensions['A'].width = 32
    ws2c.column_dimensions['B'].width = 72

    # --- 3. نقاط التكامل ---
    ws3 = wb.create_sheet("3_نقاط_التكامل", 5)
    r = sheet_header(ws3, "٣ — نقاط التكامل بين المسارين")
    headers = ["من", "إلى", "الشرط أو الحدث"]
    for c, h in enumerate(headers, 1):
        cell = ws3.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data3 = [
        ("التوظيف (Workbench)", "التدريب", "المرشح Feedback = Invalid → يُوجّه لمبيعات التدريب / التدريب"),
        ("التوظيف (تقييم المواهب)", "التدريب", "القرار = Rejected / Needs_Training → مبيعات التدريب تسوّق له الاشتراك في التدريب"),
        ("التدريب (تخرج)", "التوظيف", "بعد التخرج → الطالب يمكن أن يُعاد لمسار الترشيح أو يُسجّل Evaluated"),
        ("المالية", "كلا المسارين", "فواتير التدريب من قسم التدريب؛ إيرادات التوظيف من لوحة المالية"),
    ]
    for row in data3:
        r += 1
        for c, val in enumerate(row, 1):
            ws3.cell(r, c, val).border = border
    ws3.column_dimensions['A'].width = 28
    ws3.column_dimensions['B'].width = 14
    ws3.column_dimensions['C'].width = 72

    # --- 4. ترتيب تنفيذ الاختبار ---
    ws4 = wb.create_sheet("4_ترتيب_الاختبار", 6)
    r = sheet_header(ws4, "٤ — ترتيب تنفيذ الاختبار (اقتراح)")
    headers = ["#", "الخطوة"]
    for c, h in enumerate(headers, 1):
        cell = ws4.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    steps = [
        (1, "التسويق ومبيعات التوظيف: تسجيل مهتم بالوظيفة → تحويل لمرشح."),
        (2, "التوظيف (Workbench): فرز مرشح (Valid) ثم حجز اختبار."),
        (3, "المواهب (TA للتوظيف): إجراء الاختبار وقرار (مقبول أو مرفوض)."),
        (4, "إن كان مقبولاً: الترشيح (Matching) → مقابلة عميل → تعيين."),
        (5, "إن كان مرفوضاً أو Invalid: تحويل للتدريب → مبيعات التدريب تسوّق له → Academy → تسجيل طالب → حضور → تخرج."),
    ]
    for row in steps:
        r += 1
        ws4.cell(r, 1, row[0]).border = border
        ws4.cell(r, 2, row[1]).border = border
    ws4.column_dimensions['A'].width = 6
    ws4.column_dimensions['B'].width = 88

    # --- 5. إرشاد أول خطوة تدريب (نافذة + مستخدم) ---
    ws5 = wb.create_sheet("5_أول_خطوة_تدريب", 7)
    r = sheet_header(ws5, "٥ — إرشاد أول خطوة في اختبار التدريب")
    ws5.cell(r, 1, "الحقل").font = header_style()
    ws5.cell(r, 2, "القيمة").font = header_style()
    r += 1
    ws5.cell(r, 1, "اسم المستخدم").border = border
    ws5.cell(r, 2, "train_mgr").border = border
    r += 1
    ws5.cell(r, 1, "كلمة المرور").border = border
    ws5.cell(r, 2, "123").border = border
    r += 3
    headers5 = ["#", "ماذا تفعل", "النافذة", "المستخدم"]
    for c, h in enumerate(headers5, 1):
        cell = ws5.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    r += 1
    steps5 = [
        (1, "تسجيل الدخول", "Login", "train_mgr"),
        (2, "بعد الدخول تُوجّه إلى الأكاديمية", "Training Index (Academy)", "train_mgr"),
        (3, "إنشاء دورة إن لم توجد", "Academy — إعداد الأكاديمية — إنشاء دورة", "train_mgr"),
        (4, "إنشاء دفعة (Wave)", "Academy — جدول الدفعات — + New Wave", "train_mgr"),
        (5, "فتح تفاصيل الدفعة", "Academy — Manage أمام الدفعة", "train_mgr"),
        (6, "تسجيل طالب (يظهر لـ train_coord / trainer1)", "Academy — تسجيل طلاب → batch_details", "train_coord أو trainer1"),
    ]
    for row in steps5:
        for c, val in enumerate(row, 1):
            ws5.cell(r, c, val).border = border
        r += 1
    ws5.column_dimensions['A'].width = 6
    ws5.column_dimensions['B'].width = 38
    ws5.column_dimensions['C'].width = 45
    ws5.column_dimensions['D'].width = 24

    # --- 6أ. المسار أ — التوظيف (نافذة + مستخدم في كل صف) ---
    ws6a = wb.create_sheet("6أ_المسار_أ_التوظيف", 8)
    r = sheet_header(ws6a, "٦ — المسار أ: مسار التوظيف الكامل (من مرشح إلى تعيين)")
    headers = ["#", "العملية", "النافذة", "المستخدم", "كلمة المرور", "النتيجة المتوقعة"]
    for c, h in enumerate(headers, 1):
        cell = ws6a.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data6a = [
        (1, "تسجيل مرشح جديد", "Workbench — Quick Register", "recruiter", "123", "ظهور المرشح في الجدول"),
        (2, "إعطاء Feedback = Valid", "Workbench — Feedback أمام المرشح", "recruiter", "123", "حالة → Talent_Pool"),
        (3, "حجز اختبار", "Scheduling — Book Test — Slot وتأكيد", "recruiter", "123", "حالة → Test Scheduled"),
        (4, "إجراء التقييم (CEFR و Decision = Accepted)", "Conduct Tests — Evaluate", "ta_rec أو مستخدم Talent_Recruitment", "123", "حالة → Evaluated ثم Ready_For_Match"),
        (5, "ترشيح لطلب عميل", "Matching — Job Order ومرشح — Confirm Match", "alloc_sp أو alloc_mgr", "123", "حالة → Interview_Scheduled"),
        (6, "تسجيل نتيجة المقابلة = Accepted", "Interviews / Interview Feedback", "account أو recruiter", "123", "حالة → مقبول للتعيين"),
        (7, "إتمام التعيين", "Hiring & Onboarding — Finalize", "recruiter", "123", "Hiring Record؛ إغلاق الملف"),
    ]
    for row in data6a:
        r += 1
        for c, val in enumerate(row, 1):
            ws6a.cell(r, c, val).border = border
    ws6a.column_dimensions['A'].width = 6
    ws6a.column_dimensions['B'].width = 36
    ws6a.column_dimensions['C'].width = 42
    ws6a.column_dimensions['D'].width = 28
    ws6a.column_dimensions['E'].width = 12
    ws6a.column_dimensions['F'].width = 38

    # --- 6ب. المسار ب — التدريب (نافذة + مستخدم) ---
    ws6b = wb.create_sheet("6ب_المسار_ب_التدريب", 9)
    r = sheet_header(ws6b, "٦ — المسار ب: مسار التدريب (من أكاديمية إلى تخرج)")
    headers = ["#", "العملية", "النافذة", "المستخدم", "كلمة المرور", "النتيجة المتوقعة"]
    for c, h in enumerate(headers, 1):
        cell = ws6b.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data6b = [
        (1, "إنشاء دورة إن لم توجد", "Academy — إضافة دورة", "train_mgr", "123", "ظهور الدورة في القائمة"),
        (2, "إنشاء دفعة (Wave)", "Academy — New Wave (دورة، مدرب، قاعة، تواريخ)", "train_mgr", "123", "ظهور الدفعة في جدول Waves"),
        (3, "تسجيل طالب في الدفعة", "Academy — تسجيل طلاب (batch_details)", "train_coord أو trainer1", "123", "ظهور الطالب ضمن الدفعة"),
        (4, "تسجيل الحضور", "Attendance — اختيار الدفعة والتاريخ — مصفوفة الحضور", "trainer1", "123", "حفظ الحضور/الغياب"),
        (5, "تخريج طالب", "تفاصيل الدفعة — Graduate", "trainer1 أو manager", "123", "تحديث حالة الطالب؛ مراجعة الخريجين"),
        (6, "تسجيل قرار الطالب", "Student Decision", "من له صلاحية", "123", "تحديث السجل"),
    ]
    for row in data6b:
        r += 1
        for c, val in enumerate(row, 1):
            ws6b.cell(r, c, val).border = border
    ws6b.column_dimensions['A'].width = 6
    ws6b.column_dimensions['B'].width = 32
    ws6b.column_dimensions['C'].width = 50
    ws6b.column_dimensions['D'].width = 28
    ws6b.column_dimensions['E'].width = 12
    ws6b.column_dimensions['F'].width = 40

    # --- 6ج. المسار ج — التكامل (نافذة + مستخدم) ---
    ws6c = wb.create_sheet("6ج_المسار_ج_التكامل", 10)
    r = sheet_header(ws6c, "٦ — المسار ج: نقطة التكامل (فشل المترشح لوظيفة → تحويل للتدريب)")
    headers = ["#", "العملية", "النافذة", "المستخدم", "الهدف"]
    for c, h in enumerate(headers, 1):
        cell = ws6c.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data6c = [
        (1, "إعطاء Feedback = Invalid لمرشح", "Workbench — Feedback", "recruiter", "توجيه المرشح لمبيعات التدريب / التدريب"),
        (2, "Decision = Rejected أو Training", "Conduct Tests — Evaluate", "ta_rec أو Talent_Recruitment", "ظهور المرشح في قائمة من يحتاجون تدريباً"),
        (3, "تسجيل نفس المرشح في دفعة", "Academy — تسجيل طلاب (batch_details)", "train_coord أو trainer1", "استقبال المرشح من مسار التوظيف"),
    ]
    for row in data6c:
        r += 1
        for c, val in enumerate(row, 1):
            ws6c.cell(r, c, val).border = border
    ws6c.column_dimensions['A'].width = 6
    ws6c.column_dimensions['B'].width = 38
    ws6c.column_dimensions['C'].width = 48
    ws6c.column_dimensions['D'].width = 28
    ws6c.column_dimensions['E'].width = 52

    # --- 7. مهام لاحقة ---
    ws7 = wb.create_sheet("7_مهام_لاحقة", 11)
    r = sheet_header(ws7, "٧ — قائمة مهام للتعديل مرة واحدة لاحقاً (موثّقة)")
    headers = ["#", "المهمة", "الوصف"]
    for c, h in enumerate(headers, 1):
        cell = ws7.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data7 = [
        (1, "استنتاج أعضاء الدفعات من الشيت", "استخراج من Excel (مثلاً Guide Academy 2025) من هم في كل دفعة وعرض ذلك في النظام."),
        (2, "إظهار سجلات حضور المتدربين في ملف المتدرب", "قسم «الحضور» في صفحة ملف المتدرب يعرض سجلات الحضور/الغياب؛ استيرادها من الشيت إن وُجدت."),
    ]
    for row in data7:
        r += 1
        for c, val in enumerate(row, 1):
            ws7.cell(r, c, val).border = border
    ws7.column_dimensions['A'].width = 6
    ws7.column_dimensions['B'].width = 38
    ws7.column_dimensions['C'].width = 78

    # --- 8. سيناريو من الأسفل للأعلى (نافذة + مستخدم في كل صف) ---
    ws8 = wb.create_sheet("8_سيناريو_أسفل_لأعلى", 12)
    r = sheet_header(ws8, "٨ — سيناريو كامل من الأسفل للأعلى (من المبيعات لأعلى قسم)")
    headers = ["#", "المستوى", "العملية", "النافذة", "المستخدم", "النتيجة"]
    for c, h in enumerate(headers, 1):
        cell = ws8.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data8 = [
        (1, "١ الاستقطاب والمبيعات", "إدخال مهتم أو استيراد ليد (تعريف المهتم بالوظيفة)", "Daily Sheet / استيراد Leads", "مدير تسويق حسب الإعداد", "ظهور الليد في المصدر"),
        (2, "١ الاستقطاب والمبيعات", "تحويل الليد إلى مرشح", "لوحة المبيعات — تحويل Lead إلى Candidate", "sales", "إنشاء مرشح وربطه بالحملة"),
        (3, "٢ التوظيف", "فرز المرشح (صالح)", "Workbench — Feedback = Valid", "recruiter", "حالة → Talent_Pool"),
        (4, "٢ التوظيف", "حجز اختبار تحديد مستوى", "Scheduling — Book Test — تأكيد", "recruiter", "حالة → Test Scheduled"),
        (5, "٣ المواهب", "إجراء التقييم وإدخال النتيجة", "Conduct Tests — Evaluate — CEFR و Decision", "ta_rec أو Talent_Recruitment", "مقبول → Ready_For_Match؛ مرفوض → يُوجّه للتدريب"),
        (6, "٤ الترشيح وطلب العميل", "التأكد من وجود طلب توظيف (Job Order)", "AM Dashboard / Job Orders", "account", "وجود طلب مفتوح (أو إنشاؤه)"),
        (7, "٤ الترشيح وطلب العميل", "ترشيح المرشح لطلب العميل", "Matching — Job Order ومرشح — Confirm Match", "alloc_sp أو alloc_mgr", "حالة → Interview_Scheduled"),
        (8, "٤ الترشيح وطلب العميل", "تسجيل نتيجة مقابلة العميل (مقبول)", "Interviews / Interview Feedback = Accepted", "account أو recruiter", "جاهز لمرحلة التعيين"),
        (9, "٥ التعيين", "إتمام التعيين", "Hiring & Onboarding — Finalize", "recruiter", "سجل تعيين؛ إغلاق الملف"),
        (10, "الأعلى", "مراجعة لوحة التحكم والأقسام", "Dashboard — الهيكل الإداري، المالية، التقارير", "manager أو dev", "رؤية كل المسارات والبيانات"),
    ]
    for row in data8:
        r += 1
        for c, val in enumerate(row, 1):
            ws8.cell(r, c, val).border = border
    r += 2
    ws8.cell(r, 1, "المسار البديل — فشل المترشح لوظيفة → تحويل للتدريب (مبيعات التدريب تسوّق له الاشتراك)").font = Font(bold=True, size=12)
    r += 1
    for c, h in enumerate(headers, 1):
        cell = ws8.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    data8b = [
        ("5ب", "تدريب", "إنشاء دفعة إن لم توجد", "Academy — New Wave", "train_mgr أو train_coord", "ظهور الدفعة في الجدول"),
        ("6ب", "تدريب", "تسجيل المرشح في الدفعة", "Academy — تسجيل طلاب (batch_details)", "train_coord أو trainer1", "ظهور الطالب في الدفعة"),
        ("7ب", "تدريب", "تسجيل الحضور", "Attendance — مصفوفة الحضور", "trainer1", "حفظ حضور/غياب"),
        ("8ب", "تدريب", "تخريج الطالب", "تفاصيل الدفعة — Graduate", "trainer1 أو manager", "تحديث حالة؛ إمكانية العودة للتوظيف"),
        ("9ب", "تدريب", "(اختياري) إعادة ترشيح الخريج", "Matching أو Workbench", "recruiter أو manager", "الخريج يعود لمسار الترشيح"),
    ]
    for row in data8b:
        r += 1
        for c, val in enumerate(row, 1):
            ws8.cell(r, c, val).border = border
    ws8.column_dimensions['A'].width = 8
    ws8.column_dimensions['B'].width = 20
    ws8.column_dimensions['C'].width = 42
    ws8.column_dimensions['D'].width = 42
    ws8.column_dimensions['E'].width = 28
    ws8.column_dimensions['F'].width = 40

    # --- 9. ملخص اليوزرات (من app.py / قاعدة البيانات) ---
    ws9 = wb.create_sheet("9_اليوزرات", 13)
    r = sheet_header(ws9, "ملخص اليوزرات في النظام (كلمة المرور: 123 — من app.py / قاعدة البيانات)")
    headers = ["الدور", "اليوزر", "ملاحظة"]
    for c, h in enumerate(headers, 1):
        cell = ws9.cell(r, c, h)
        cell.font = header_style()
        cell.fill = header_fill
        cell.border = border
    users = [
        ("مدير عام", "manager", "app.py init"),
        ("مطور", "dev", "app.py init"),
        ("مبيعات (توظيف)", "sales", "app.py init"),
        ("توظيف (Recruiter)", "recruiter", "app.py init"),
        ("مدير توظيف", "rec_mgr", "app.py init"),
        ("مدير حسابات", "account", "app.py init"),
        ("مسكّن / تخصيص", "alloc_sp أو alloc_mgr", "app.py init"),
        ("تقييم مواهب توظيف (Talent_Recruitment)", "ta_rec", "create_users إن وُجد"),
        ("مدير تدريب", "train_mgr", "app.py training_users"),
        ("رئيس قسم تدريب", "train_head", "app.py training_users"),
        ("قائد تدريب", "train_lead", "app.py training_users"),
        ("منسق تدريب", "train_coord", "app.py training_users"),
        ("مبيعات التدريب", "train_sales", "app.py training_users"),
        ("مختبر مواهب التدريب (TA-Training)", "ta_train", "create_users إن وُجد"),
        ("مدرب", "trainer1", "app.py training_users / init"),
    ]
    for row in users:
        r += 1
        for c, val in enumerate(row, 1):
            ws9.cell(r, c, val).border = border
    ws9.column_dimensions['A'].width = 36
    ws9.column_dimensions['B'].width = 28
    ws9.column_dimensions['C'].width = 24

    try:
        wb.save(OUTPUT)
        print(f"تم إنشاء الملف: {OUTPUT}")
    except PermissionError:
        alt = "خريطة_اختبار_تكامل_النظام_جديد.xlsx"
        wb.save(alt)
        print(f"الملف الأصلي مفتوح. تم الحفظ باسم: {alt}")

if __name__ == "__main__":
    main()
