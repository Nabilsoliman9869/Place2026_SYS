# -*- coding: utf-8 -*-
"""
استنتاج مختبري مواهب التدريب (Talent_Training) من ملف Guide Academy Placements 2025.xlsx
ثم إنشاء المستخدمين ومواعيد TASchedules حتى تظهر «الشاغر المتاحة» في حجز اختبار التدريب.

شغّل: python infer_ta_training_from_excel.py
يحتاج: openpyxl, pyodbc, db_config.json، والملف Guide Academy Placements 2025.xlsx في نفس المجلد.
"""
import os
import sys
import json
import re
from datetime import datetime, timedelta

try:
    import openpyxl
except ImportError:
    print("ثبّت openpyxl: pip install openpyxl")
    sys.exit(1)
try:
    import pyodbc
except ImportError:
    print("ثبّت pyodbc: pip install pyodbc")
    sys.exit(1)

BASE = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE, "Guide Academy Placements 2025.xlsx")


def safe_str(val, max_len=200):
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() in ("none", "nan", "#n/a", "-"):
        return None
    return s[:max_len] if len(s) > max_len else s


def name_to_username(full_name):
    """تحويل الاسم إلى username (حروف إنجليزية صغيرة، بدون مسافات)."""
    if not full_name:
        return None
    s = str(full_name).strip()
    # إزالة أي حروف غير لاتينية أو أرقام أو مسافات، واستبدال المسافات بـ _
    parts = re.split(r"\s+", s)
    if not parts:
        return None
    # نأخذ أول جزء (الاسم الأول عادة) أو نربط الأجزاء
    base = "".join(c for c in parts[0] if c.isalnum()).lower() or "ta"
    if not base:
        return None
    # إن كان الاسم عربي أو فيه رموز، نستخدم أول حرف أو نولد من الاسم
    if not re.match(r"^[a-z0-9]+$", base):
        base = "ta_" + str(abs(hash(s)) % 10000)
    return base[:50]


def row_to_dict(header, row):
    d = {}
    for i, h in enumerate(header):
        if h and i < len(row):
            key = str(h).strip()
            if key and key != "@dropdown":
                d[key] = row[i]
    return d


def load_sheet_interviewers(wb, sheet_name):
    """استخراج قائمة أسماء المقابلين/المختبِرين من ورقة (أي عمود اسمه Interviewer أو GATB Interviewer)."""
    if sheet_name not in wb.sheetnames:
        return set()
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return set()
    header = [str(x).strip() if x else "" for x in rows[0]]
    names = set()
    for r in rows[1:]:
        if not r:
            continue
        d = row_to_dict(rows[0], r)
        for col in ("Interviewer", "GATB Interviewer", "GATB Interviewer "):
            v = safe_str(d.get(col))
            if v:
                names.add(v)
    return names


def get_all_ta_training_names_from_excel():
    """جمع كل الأسماء الفريدة لمختبري/مقابلي التدريب من الأوراق ذات الصلة."""
    if not os.path.exists(EXCEL_PATH):
        print(f"الملف غير موجود: {EXCEL_PATH}")
        return set()
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    all_names = set()
    for sheet in ("GA Interviews", "Exit Make-Up", "Foundation(certain program)", "train to hire (certain program)"):
        names = load_sheet_interviewers(wb, sheet)
        all_names |= names
        if names:
            print(f"  من ورقة '{sheet}': {len(names)} اسم — {list(names)[:5]}{'...' if len(names) > 5 else ''}")
    wb.close()
    return all_names


def get_db_connection():
    config_path = os.path.join(BASE, "db_config.json")
    if not os.path.exists(config_path):
        raise FileNotFoundError("لا يوجد db_config.json")
    with open(config_path, "r", encoding="utf-8") as f:
        config = json.load(f)
    server = config.get("server", ".")
    port = config.get("port", "1433")
    database = config.get("database", "Place2026DB")
    username = config.get("username", "")
    password = config.get("password", "")
    use_trusted = config.get("use_trusted", False)
    conn_str = (
        f"DRIVER={{ODBC Driver 17 for SQL Server}};SERVER={server},{port};DATABASE={database};"
        + ("Trusted_Connection=yes;" if use_trusted else f"UID={username};PWD={password};")
    )
    return pyodbc.connect(conn_str, timeout=15)


def ensure_user(conn, full_name, role="Talent_Training", password="123"):
    """إنشاء مستخدم إن لم يكن موجوداً، أو تحديث الرول إلى Talent_Training."""
    cursor = conn.cursor()
    username = name_to_username(full_name)
    if not username:
        return None
    # تجنب تكرار username: إن وُجد مستخدم بنفس الاسم الكامل نستخدمه
    cursor.execute("SELECT UserID, Role FROM Users_1 WHERE FullName = ? OR Username = ?", (full_name, username))
    row = cursor.fetchone()
    if row:
        user_id, current_role = row.UserID, row.Role
        if current_role != role:
            cursor.execute("UPDATE Users_1 SET Role = ? WHERE UserID = ?", (role, user_id))
            print(f"    تحديث الرول لـ {full_name} -> {role}")
        return user_id
    # إنشاء مستخدم جديد (إن كان username مستخدماً نضيف لاحقة)
    base_username = username
    suffix = 0
    while True:
        try_username = base_username if suffix == 0 else f"{base_username}_{suffix}"
        cursor.execute("SELECT UserID FROM Users_1 WHERE Username = ?", (try_username,))
        if cursor.fetchone() is None:
            break
        suffix += 1
    cursor.execute(
        "INSERT INTO Users_1 (Username, Password, Role, FullName) VALUES (?, ?, ?, ?)",
        (try_username, password, role, full_name),
    )
    conn.commit()
    cursor.execute("SELECT SCOPE_IDENTITY()")
    user_id = int(cursor.fetchone()[0])
    print(f"    إضافة مستخدم: {try_username} ({full_name}) -> {role}")
    return user_id


def slots_10am_to_9pm_every_15min():
    """مواعيد كل 15 دقيقة من 10 صباحاً إلى 9 مساءً (10:00 — 21:00)."""
    times = []
    for hour in range(10, 22):
        for minute in (0, 15, 30, 45):
            if hour == 21 and minute != 0:
                continue
            times.append(f"{hour:02d}:{minute:02d}")
    return times  # 44 موعد (10:00–20:45 كل 15 دقيقة) + 21:00 = 45 موعد/يوم


def ensure_taschedules_slots(conn, user_id, days=14, slot_times=None):
    """إنشاء مواعيد TASchedules لمختبر (EvaluatorID): كل 15 دقيقة من 10 ص إلى 9 م."""
    cursor = conn.cursor()
    start = datetime.today().date()
    if slot_times is None:
        slot_times = slots_10am_to_9pm_every_15min()
    added = 0
    for d in range(days):
        slot_date = (start + timedelta(days=d)).strftime("%Y-%m-%d")
        for slot_time in slot_times:
            cursor.execute(
                "SELECT SlotID FROM TASchedules WHERE EvaluatorID = ? AND SlotDate = ? AND SlotTime = ?",
                (user_id, slot_date, slot_time),
            )
            if cursor.fetchone() is None:
                cursor.execute(
                    "INSERT INTO TASchedules (SlotDate, SlotTime, Status, EvaluatorID) VALUES (?, ?, 'Available', ?)",
                    (slot_date, slot_time, user_id),
                )
                added += 1
    conn.commit()
    return added


# ما يفعله مختبرو مواهب التدريب (Talent_Training) ونوافذهم في النظام
TA_TRAINING_ROLE_DESC = """
مختبرو مواهب التدريب (Talent_Training / TA-Training):
• المهمة: إجراء اختبار المواهب للتدريب (تقييم لغوي/مهاري للمهتمين بالتدريب قبل الالتحاق بالدورة).
• النوافذ: لوحة مواعيد اختبار التدريب (talent_dashboard)، تنفيذ الاختبار وتدوين النتيجة (conduct test)، متابعة المواعيد المحجوزة لهم في TASchedules.
• من يسجل الموعد: مبيعات التدريب (train_sales) أو المدير من صفحة «جدولة اختبار التدريب» ثم «حجز موعد» — يختار شاغراً من المواعيد المتاحة لأي مختبر Talent_Training.
"""


def write_names_and_windows_doc(names, out_path=None):
    """توثيق الأسماء المستنتجة + مهامهم ونوافذهم."""
    out_path = out_path or os.path.join(BASE, "مختبرو_مواهب_التدريب_من_الاكسل.md")
    lines = [
        "# مختبرو مواهب التدريب المستنتجون من Guide Academy Placements 2025.xlsx",
        "",
        "## الأسماء المستنتجة (من أعمدة Interviewer / GATB Interviewer في أوراق GA، Exit، Foundation، train to hire)",
        "",
    ]
    for n in sorted(names):
        lines.append(f"- {n}")
    lines.extend(["", "---", "", TA_TRAINING_ROLE_DESC.strip(), ""])
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"\nتم حفظ التوثيق: {out_path}")


def main():
    print("=" * 60)
    print("استنتاج مختبري مواهب التدريب من Guide Academy Placements 2025.xlsx")
    print("=" * 60)
    names = get_all_ta_training_names_from_excel()
    if not names:
        print("لم يُعثر على أسماء مختبِرين في الملف. تأكد من وجود أعمدة 'Interviewer' أو 'GATB Interviewer' في أوراق GA / Exit.")
        return
    print(f"\nإجمالي أسماء فريدة: {len(names)}")
    write_names_and_windows_doc(names)
    conn = get_db_connection()
    user_ids = []
    for full_name in sorted(names):
        uid = ensure_user(conn, full_name)
        if uid:
            user_ids.append(uid)
    if not user_ids:
        print("لم يُضف أي مستخدم جديد. قد يكونون موجودين مسبقاً.")
    else:
        print(f"\nتم التأكد من {len(user_ids)} مستخدم بدور Talent_Training.")
    # إنشاء مواعيد لجميع مستخدمي Talent_Training (بما فيهم الموجودون سابقاً)
    cursor = conn.cursor()
    cursor.execute("SELECT UserID, FullName, Username FROM Users_1 WHERE Role IN ('Talent_Training', 'TA-Training')")
    ta_users = cursor.fetchall()
    conn.close()
    if not ta_users:
        print("لا يوجد مستخدمون بدور Talent_Training في القاعدة. أضفهم يدوياً أو راجع السكربت.")
        return
    slot_times = slots_10am_to_9pm_every_15min()
    print(f"\nإنشاء مواعيد TASchedules لـ {len(ta_users)} مختبِر: كل 15 دقيقة من 10 ص إلى 9 م ({len(slot_times)} موعد/يوم)، لمدة 14 يوماً.")
    conn = get_db_connection()
    total_slots = 0
    for u in ta_users:
        n = ensure_taschedules_slots(conn, u.UserID, days=14, slot_times=slot_times)
        total_slots += n
        if n:
            print(f"  {u.FullName or u.Username}: +{n} موعد")
    conn.close()
    print(f"\nتم. إجمالي مواعيد جديدة: {total_slots}. جرّب الآن «حجز موعد اختبار مواهب تدريب» من مبيعات التدريب.")


if __name__ == "__main__":
    main()
